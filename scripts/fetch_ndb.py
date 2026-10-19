"""
Fetch Notifiable Data Breaches (NDB) statistics from the OAIC's official
dataset on data.gov.au.

The OAIC publishes each reporting period as an XLSX resource on the
"notifiable-data-breaches-ndb-scheme" dataset. The resource is discovered via
the data.gov.au CKAN API, so new periods are picked up automatically.
"""
import io
import re
from datetime import datetime, timezone

import openpyxl

from .config import SOURCES
from .utils import logger, save_data, fetch_url


DATASET_ID = "notifiable-data-breaches-ndb-scheme"
DATASET_PAGE = f"https://data.gov.au/data/dataset/{DATASET_ID}"

# Breach-source labels the OAIC uses as sub-rows under months and sectors
SOURCE_LABELS = {
    "Malicious or criminal attack", "Human error", "System fault",
    "Other", "Currently unknown",
}
# Top-level categories in the "Source of breach" sheet
SOURCE_CATEGORIES = {"Malicious or criminal attack", "Human error", "System fault"}


def _latest_xlsx_resource() -> dict | None:
    """Find the most recently published XLSX resource on the OAIC dataset."""
    resp = fetch_url(SOURCES["datagovau_package_show"], params={"id": DATASET_ID})
    if not resp:
        return None
    resources = resp.json().get("result", {}).get("resources", [])
    xlsx = [r for r in resources if (r.get("format") or "").upper() == "XLSX"]
    if not xlsx:
        logger.warning("  No XLSX resource on the OAIC NDB dataset")
        return None
    return max(xlsx, key=lambda r: r.get("last_modified") or r.get("created") or "")


def _sheet(wb, name: str):
    """Look up a worksheet by name, ignoring case and stray whitespace."""
    for ws in wb.worksheets:
        if ws.title.strip().lower() == name.lower():
            return ws
    return None


def _label_rows(ws, stop_at_total: bool = True):
    """Yield (label, count) pairs from the first two columns of a sheet."""
    for row in ws.iter_rows(min_col=1, max_col=2, values_only=True):
        label, value = row[0], row[1] if len(row) > 1 else None
        if not isinstance(label, str) or not isinstance(value, (int, float)):
            continue
        label = label.strip()
        if label == "Grand Total":
            if stop_at_total:
                return
            continue
        yield label, int(value)


def _sheet_title(ws) -> str:
    first = next(ws.iter_rows(min_row=1, max_row=1, max_col=1, values_only=True), (None,))
    return first[0] if isinstance(first[0], str) else ""


def _period_from_title(title: str) -> str:
    """Pull "July-December 2025" from a sheet title like '..., July-December 2025'."""
    match = re.search(r"([A-Z][a-z]+\s*-\s*[A-Z][a-z]+\s+\d{4})\s*$", title.strip())
    return re.sub(r"\s*-\s*", "-", match.group(1)) if match else ""


def _grouped(ws, child_labels: set) -> list[dict]:
    """Parse a sheet of parent rows each followed by breach-source sub-rows."""
    groups = []
    for label, count in _label_rows(ws):
        if label in child_labels and groups:
            groups[-1]["by_source"][label] = count
        else:
            groups.append({"name": label, "count": count, "by_source": {}})
    return groups


def _specific_sources(ws) -> list[dict]:
    """Parse the 'Source of breach' sheet into category/source rows."""
    rows, category, seen = [], None, set()
    for label, count in _label_rows(ws, stop_at_total=False):
        if label in SOURCE_CATEGORIES:
            category = label
            continue
        if category and (category, label) not in seen:
            seen.add((category, label))
            rows.append({"category": category, "source": label, "count": count})
    return rows


def parse_workbook(content: bytes) -> dict:
    """Extract NDB statistics from an OAIC NDB XLSX file.

    Each sheet is parsed independently so a layout change in one sheet
    doesn't lose the rest.
    """
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    out = {}

    def attempt(key, sheet_name, fn):
        ws = _sheet(wb, sheet_name)
        if ws is None:
            logger.warning(f"  NDB sheet '{sheet_name}' not found")
            return
        try:
            out[key] = fn(ws)
            if "period" not in out:
                out["period"] = _period_from_title(_sheet_title(ws))
        except Exception as e:
            logger.error(f"  Could not parse NDB sheet '{sheet_name}': {e}")

    attempt("months", "NDB by month", lambda ws: _grouped(ws, SOURCE_LABELS))
    attempt("sectors", "Top 5 sectors by source", lambda ws: _grouped(ws, SOURCE_LABELS))
    attempt("specific_sources", "Source of breach", _specific_sources)
    attempt("personal_information", "Personal information",
            lambda ws: [{"kind": k, "count": c} for k, c in _label_rows(ws)])
    attempt("individuals_affected", "Individuals affected",
            lambda ws: [{"range": r, "count": c} for r, c in _label_rows(ws)])
    return out


def run():
    """Fetch the latest OAIC NDB dataset and save summary statistics."""
    logger.info("Fetching OAIC NDB statistics from data.gov.au...")

    resource = _latest_xlsx_resource()
    if not resource:
        raise RuntimeError("OAIC NDB dataset unavailable")

    resp = fetch_url(resource["url"], timeout=60)
    if not resp:
        raise RuntimeError("Could not download OAIC NDB spreadsheet")

    parsed = parse_workbook(resp.content)
    months = parsed.get("months", [])

    # Totals by breach source across the whole period
    by_source = {}
    for m in months:
        for src, n in m["by_source"].items():
            by_source[src] = by_source.get(src, 0) + n

    trend = [{
        "period": m["name"],
        "total": m["count"],
        "malicious": m["by_source"].get("Malicious or criminal attack", 0),
        "human_error": m["by_source"].get("Human error", 0),
        "system_faults": m["by_source"].get("System fault", 0),
        "other": m["by_source"].get("Other", 0) + m["by_source"].get("Currently unknown", 0),
    } for m in months]

    data = {
        "fetched_at": datetime.now(timezone.utc).isoformat(),
        "source": "OAIC Notifiable Data Breaches dataset (data.gov.au)",
        "source_url": DATASET_PAGE,
        "resource_name": resource.get("name", ""),
        "resource_url": resource.get("url", ""),
        "resource_last_modified": resource.get("last_modified", ""),
        "period": parsed.get("period") or resource.get("name", ""),
        "total": sum(m["count"] for m in months),
        "by_source": by_source,
        "trend_summary": trend,
        "sectors": [{"sector": s["name"], "count": s["count"], "by_source": s["by_source"]}
                    for s in parsed.get("sectors", [])],
        "specific_sources": parsed.get("specific_sources", []),
        "personal_information": parsed.get("personal_information", []),
        "individuals_affected": parsed.get("individuals_affected", []),
    }
    logger.info(f"  OAIC NDB: {data['period']}, {data['total']} notifications")

    save_data("ndb.json", data)
    return data


if __name__ == "__main__":
    run()
